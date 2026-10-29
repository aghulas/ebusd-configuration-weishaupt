#!/usr/bin/env python3
"""Build the Weishaupt catalog (catalog/*.csv) from an export of the "WCM" workbook.

Usage: build_catalog.py WCM.xlsx

The workbook (Google Sheets "WCM", exported as .xlsx) holds the WTC memory map named after the .SYC symbol
tables, the EEPROM dumps of both burners with the parameter categories, the regulator registers (0902/0903) and
the ComfortSoft 5022 info values. This script only reorganises that material into one row per point; it never
talks to the bus. Requires openpyxl.
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "catalog"

EEPROM_CATEGORIES = {7: "manufacturer", 8: "heating_specialist", 9: "user", 10: "errors", 11: "statistics",
                     12: "ion_calibration"}
LEVELS = {"manufacturer": "manufacturer", "heating_specialist": "heating_specialist", "user": "user"}

# EEPROM blocks protected by a check byte (first byte, last byte, check byte), see block_check() and README.
# 0x053-0x089 and 0x08C-0x0A2 are redundant copies of the heating specialist and manufacturer blocks.
BLOCK_CRCS = [(0x002, 0x017, 0x018), (0x01B, 0x050, 0x051), (0x053, 0x088, 0x089), (0x08C, 0x0A1, 0x0A2),
              (0x1EE, 0x1F5, 0x1F6), (0x300, 0x303, 0x304)]


def block_check(data: bytes) -> int:
    """Check byte of an EEPROM block, as computed by WCM-Diagnose (function 0x64D5FC of WCMDiag 5.5.19)."""
    value = 0
    for byte in data:
        value <<= 1
        if value > 0xFF:
            value = (value - 0x100) ^ 0x21
        value ^= byte
    return 0xFF - value

# WCM-Diagnose parameters matched by symbol name and by the value displayed on 29/10/2026 (bc1)
PARAMS = {
    "V24_KM_ADR": ("P12", "Adresse de la chaudière (A=1, B=2)", "", "", "name+value (bc1=1/A, bc2=2/B)"),
    "SONDERNIVEAU": ("P18", "Niveau spécial", "°C", "", "name+value"),
    "TAUSS_KORREKTUR": ("P20", "Correction sonde extérieure", "K", "", "name+value"),
    "MINTVSOLL": ("P30", "Température de départ minimale", "°C", "", "name+value"),
    "MAXTVSOLL": ("P31", "Température de départ maximale", "°C", "", "name+value"),
    "MXTABGAS": ("P33", "Température de coupure fumées", "°C", "", "name+value"),
    "TBRMINOFF": ("P34", "Temps de blocage brûleur", "min", "", "name+value"),
    "NG_ZUEND": ("P35", "Charge d'allumage", "%", "", "name+value"),
    "NG_HZ_MIN": ("P36", "Puissance minimale chauffage", "%", "", "name+value (23 % aussi NG_HZ_FREI, NG_WW_MIN)"),
    "NG_HZ_MAX": ("P37", "Puissance maximale chauffage", "%", "", "name+value"),
    "NG_WW_MAX": ("P38", "Puissance maximale ECS", "%", "", "name+value"),
    "EEA_TWWSOLL_EPM": ("", "Consigne ECS", "°C", "", "name+value (50 °C)"),
    "EEA_TVSOLL_EPM": ("", "Consigne de départ", "°C", "", "name"),
    "EEA_TVABSENK": ("", "Abaissement de la consigne de départ", "", "", "name"),
    "EEA_WISO_EPM": ("", "Été / hiver", "", "", "name"),
}


def crc_weishaupt(payload: bytes) -> int:
    crc = payload[0]
    for byte in payload[1:]:
        for _ in range(8):
            crc = ((crc << 1) ^ 0x5C) & 0xFF if crc & 0x80 else (crc << 1) & 0xFF
        crc ^= byte
    return crc


def read_id(space: str, address: int, size: int) -> str:
    """ID of a 5000 read (CRC included), as in WH11928.inc for one byte and wtc.history.inc for longer reads."""
    lo = address & 0xFF
    if space == "eeprom" and address >= 0x100:
        payload = bytes([0x06, address >> 8, ((size - 1) << 4) | 0x02, lo])
    else:
        mem_type = {"ram": 0x01, "eeprom": 0x02, "xram": 0x03}[space]
        payload = bytes([((size - 1) << 4) | mem_type, lo])
    return f"{crc_weishaupt(payload):02x}{payload.hex()}"


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def block_crc(address: int) -> str:
    for lo, hi, crc_address in BLOCK_CRCS:
        if lo <= address <= hi:
            return f"0x{crc_address:03x}"
    return ""


def wtc_rows(workbook):
    rows = []
    # EEPROM: pages 0600-0603, one row per byte: page, address in page, -, hex, name, "abs value" bc1/bc2, categories
    for row in list(workbook["EEPROM"].iter_rows(values_only=True))[2:]:
        if not row or row[0] is None:
            continue
        page = int(str(row[0]).strip('"')) - 600
        address = page * 0x100 + int(row[1])
        name = text(row[4])
        categories = [EEPROM_CATEGORIES[i] for i in EEPROM_CATEGORIES if i < len(row) and text(row[i])]
        # a byte shown at several levels is reachable from the least privileged one
        category = next((c for c in ("user", "heating_specialist", "manufacturer") if c in categories),
                        categories[0] if categories else "")
        refs = [text(row[i]).split()[-1] if text(row[i]) else "" for i in (5, 6)]
        if not name and not category:
            continue
        rows.append({"space": "eeprom", "address": address, "name": name, "category": category, "ref_bc1": refs[0],
                     "ref_bc2": refs[1], "ref_value": "", "notes": ""})
    # RAM and XRAM: address, -, hex, name, -, -, value of one dump, annotations under Config/Live
    for sheet, space, base in (("RAM", "ram", 0x0000), ("X-RAM", "xram", 0xF000)):
        table = list(workbook[sheet].iter_rows(values_only=True))
        header = {i: text(c).lower() for i, c in enumerate(table[0]) if i >= 7 and text(c)}
        for row in table[2:]:
            if not row or row[0] is None:
                continue
            name = text(row[3])
            marks = {header.get(i, ""): text(row[i]) for i in range(7, min(len(row), 10)) if text(row[i])}
            notes = "; ".join(v for v in marks.values() if not re.fullmatch(r"[\d.]+", v))
            if not name and not notes:
                continue
            rows.append({"space": space, "address": base + int(row[0]), "name": name,
                         "category": next((k for k in marks if k), ""), "ref_bc1": "", "ref_bc2": "",
                         "ref_value": text(row[6]), "notes": notes})
    # sizes: fault counters are 16 bits, statistics counters span up to the next named byte (max 3)
    rows.sort(key=lambda r: ({"ram": 0, "xram": 1, "eeprom": 2}[r["space"]], r["address"]))
    for i, r in enumerate(rows):
        size = 1
        if r["name"].startswith("EEA_FC_"):
            size = 2
        elif r["category"] == "statistics" and r["name"]:
            nxt = next((n["address"] for n in rows[i + 1:] if n["space"] == r["space"] and n["name"]), None)
            if nxt is not None and 1 < nxt - r["address"] <= 3:
                size = nxt - r["address"]
        r["size"] = size
    out = []
    for r in rows:
        # parameters live in the EEPROM blocks; RAM holds working copies under the same names
        mapped = r["space"] == "eeprom" and (0x01B <= r["address"] <= 0x050 or 0x300 <= r["address"] <= 0x303)
        param, label, unit, divider, confidence = PARAMS.get(r["name"], ("",) * 5) if mapped else ("",) * 5
        out.append({
            "space": r["space"], "address": f"0x{r['address']:04x}", "name": r["name"], "size": r["size"],
            "category": r["category"], "level": LEVELS.get(r["category"], "info"),
            "block_crc": block_crc(r["address"]) if r["space"] == "eeprom" else "",
            "param": param, "label_fr": label, "unit": unit, "divider": divider, "confidence": confidence,
            "read_id": read_id(r["space"], r["address"], r["size"]),
            "ref_bc1": r["ref_bc1"], "ref_bc2": r["ref_bc2"], "ref_value": r["ref_value"], "notes": r["notes"],
        })
    return out


def regulator_rows(workbook, register_inc: Path):
    defined = {}
    for line in register_inc.read_text(encoding="utf-8", errors="replace").splitlines():
        parts = line.split(",")
        if len(parts) > 7 and parts[0] in ("r", "w") and parts[7].strip('"'):
            # read IDs are register, bank, length (6 digits); write IDs only register and bank
            defined.setdefault(parts[7].strip('"').upper()[:4], set()).add(parts[0])
    out = []
    for row in list(workbook["09020903"].iter_rows(values_only=True))[1:]:
        row = list(row) + [None] * (6 - len(row))
        name, comment = text(row[0]), text(row[1])
        register, bank = text(row[3]).upper(), text(row[4]).upper()
        if not name or not register:
            continue
        menu = re.search(r"\((?:User\s*)?(\d+)\)", comment)
        templates = [text(c) for c in row[8:] if text(c).startswith("_")]
        ident = f"{register.zfill(2)}{bank.zfill(2)}02"
        out.append({"name": name, "comment": comment, "menu": menu.group(1) if menu else "",
                    "register": register.zfill(2), "bank": bank.zfill(2), "id": ident,
                    "templates": " ".join(templates),
                    "register_inc": "".join(sorted(defined.get(ident.upper()[:4], set())))})
    return out


def info_5022_rows(workbook):
    out = []
    for row in workbook["5022"].iter_rows(values_only=True):
        row = list(row) + [None] * (6 - len(row))
        name = text(row[0])
        if not name:
            continue
        cells = [text(c).rstrip(",").strip() for c in row]
        numbers = [c for c in cells[5:] if re.fullmatch(r"-?[\d.]+", c)]
        flags = [c for c in cells[5:] if c in ("True", "False")]
        unit = next((c for c in cells[5:] if c and not re.fullmatch(r"-?[\d.]+|True|False", c)), "")
        out.append({"name": name, "zz": cells[1], "pbsb": cells[2].split(".")[0], "id": cells[4],
                    "factor": numbers[0] if numbers else "", "flags": " ".join(flags), "unit": unit})
    return out


def write(path: Path, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    print(f"{path.relative_to(OUT.parent)}: {len(rows)} rows")


def main():
    workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    OUT.mkdir(exist_ok=True)
    write(OUT / "wtc-memory.csv", wtc_rows(workbook))
    write(OUT / "regulator-registers.csv",
          regulator_rows(workbook, OUT.parent / "weishaupt" / "register.inc"))
    write(OUT / "wtc-info-5022.csv", info_5022_rows(workbook))


if __name__ == "__main__":
    main()
